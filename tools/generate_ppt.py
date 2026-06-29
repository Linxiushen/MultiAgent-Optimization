# -*- coding: utf-8 -*-
"""
多Agent智能分流 + 长程记忆优化系统 —— 演示文稿生成器

使用 python-pptx 生成一套「简约浅色商务风」的可编辑 .pptx。
所有内容取自项目 README / TECHNICAL_HIGHLIGHTS / PERFORMANCE_BENCHMARKS。

运行:
    python tools/generate_ppt.py
产物:
    MultiAgent_Optimization.pptx  (16:9, 可用 PowerPoint / WPS / Keynote 打开再编辑)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- 设计 token
# 简约浅色商务风：白底 + 一个克制的主色 + 中性灰，大量留白
INK = RGBColor(0x1F, 0x2A, 0x37)     # 主文字 近黑
SUB = RGBColor(0x5B, 0x66, 0x72)     # 次要文字 灰
HAIR = RGBColor(0xE3, 0xE8, 0xEE)    # 分隔线 浅灰
BG = RGBColor(0xFF, 0xFF, 0xFF)      # 背景 纯白
SOFT = RGBColor(0xF5, 0xF7, 0xFA)    # 卡片底 极浅灰
BRAND = RGBColor(0x2D, 0x6C, 0xDF)   # 主色 商务蓝
BRAND_D = RGBColor(0x1E, 0x4F, 0xA8) # 主色深
ACCENT = RGBColor(0x12, 0xB7, 0x8F)  # 点缀 青绿
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Microsoft YaHei"   # 中文回退；缺失时 WPS/Office 会自动替换
FONT_NUM = "Arial"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)


def _set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _no_line(shape):
    shape.line.fill.background()


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    _no_line(shape)


def _shadow_off(shape):
    # python-pptx 默认 autoshape 会带预设阴影，关掉保持扁平
    sp = shape.shape_element if hasattr(shape, "shape_element") else shape._element
    spPr = sp.find(qn("p:spPr"))
    if spPr is None:
        return
    # 移除已有 effectLst 并加空的
    for tag in ("a:effectLst",):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.0, space_after=2):
    """runs: list of (text, size, color, bold) ；每个 run 单独成段或同段。
    这里简化为每个元素是一段。"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, (text, size, color, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = FONT
    return tb


def add_rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    _solid(sp, color)
    _shadow_off(sp)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def add_line(slide, x, y, w, color=HAIR, weight=1.0):
    ln = slide.shapes.add_connector(2, x, y, x + w, y)
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def kicker(slide, x, y, text, color=BRAND):
    """小标签 + 短横线，页眉的栏目感"""
    bar = add_rect(slide, x, y + Emu(int(Pt(2))), Pt(18), Pt(4), color)
    add_text(slide, x + Pt(26), y - Pt(4), Inches(6), Pt(24),
             [(text, 13, color, True)])


def footer(slide, page):
    add_line(slide, Inches(0.7), Inches(7.0), Inches(11.93), HAIR, 0.75)
    add_text(slide, Inches(0.7), Inches(7.05), Inches(8), Pt(18),
             [("多Agent智能分流 + 长程记忆优化系统", 9, SUB, False)])
    add_text(slide, Inches(11.6), Inches(7.05), Inches(1.0), Pt(18),
             [(f"{page:02d}", 9, SUB, True)], align=PP_ALIGN.RIGHT)


prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


def new_slide():
    s = prs.slides.add_slide(BLANK)
    _set_bg(s)
    return s


def title_block(slide, kick, title, sub=None):
    kicker(slide, Inches(0.7), Inches(0.62), kick)
    add_text(slide, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.9),
             [(title, 30, INK, True)])
    if sub:
        add_text(slide, Inches(0.7), Inches(1.62), Inches(11.9), Inches(0.5),
                 [(sub, 13, SUB, False)])


# ================================================================ 1. 封面
s = new_slide()
# 左侧细色条
add_rect(s, 0, 0, Pt(10), EMU_H, BRAND)
# 顶部品牌行
add_text(s, Inches(0.9), Inches(0.7), Inches(8), Pt(20),
         [("MULTI-AGENT  ·  LONG-TERM MEMORY", 11, SUB, True)])
# 主标题
add_text(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.4),
         [("多Agent智能分流", 46, INK, True),
          ("＋长程记忆优化系统", 46, BRAND, True)],
         line_spacing=1.05)
add_line(s, Inches(0.95), Inches(4.35), Inches(2.2), BRAND, 2.5)
add_text(s, Inches(0.9), Inches(4.55), Inches(10.5), Inches(0.9),
         [("智能路由 · 长程记忆 · 分布式协调 · 故障恢复 · 响应聚合", 15, SUB, False)])
# 底部信息
add_text(s, Inches(0.9), Inches(6.55), Inches(11), Pt(20),
         [("技术方案演示  |  Python · FAISS · PyTorch · Sentence-Transformers", 11, SUB, False)])

# ================================================================ 2. 项目概述
s = new_slide()
title_block(s, "OVERVIEW", "项目概述",
            "基于多Agent架构的智能分流系统，结合长程记忆优化，提升AI系统响应质量与效率")
intro = ("系统通过智能路由将用户查询分配给最适合的Agent处理，利用长程记忆存储与检索"
         "相关信息以提供更准确、个性化的回答；同时具备分布式协调与恢复管理能力，"
         "确保高可用性与容错性。")
add_text(s, Inches(0.7), Inches(2.25), Inches(11.9), Inches(0.9),
         [(intro, 14, INK, False)], line_spacing=1.35)

cards = [
    ("智能分流", "按查询类型路由到最合适的 Agent"),
    ("长程记忆", "FAISS 向量检索 + 记忆优化压缩"),
    ("分布式协调", "注册发现 · 负载均衡 · 资源锁"),
    ("高可用恢复", "心跳监控 · 检查点 · 平滑降级"),
]
cw, gap = Inches(2.85), Inches(0.18)
x0 = Inches(0.7)
y0 = Inches(3.5)
for i, (h, d) in enumerate(cards):
    x = x0 + i * (cw + gap)
    card = add_rect(s, x, y0, cw, Inches(2.3), SOFT,
                    MSO_SHAPE.ROUNDED_RECTANGLE, 0.06)
    add_rect(s, x + Inches(0.25), y0 + Inches(0.28), Pt(34), Pt(4), BRAND)
    add_text(s, x + Inches(0.25), y0 + Inches(0.5), cw - Inches(0.5), Pt(30),
             [(h, 17, INK, True)])
    add_text(s, x + Inches(0.25), y0 + Inches(1.05), cw - Inches(0.5), Inches(1.1),
             [(d, 12, SUB, False)], line_spacing=1.3)
footer(s, 2)

# ================================================================ 3. 系统架构
s = new_slide()
title_block(s, "ARCHITECTURE", "系统架构",
            "查询自上而下流经路由、Agent群、记忆与聚合层，由分布式协调与恢复层横向支撑")

def arch_box(slide, x, y, w, h, text, fill, txt=WHITE, sub=None, size=14):
    b = add_rect(slide, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    runs = [(text, size, txt, True)]
    if sub:
        runs.append((sub, 9.5, txt, False))
    add_text(slide, x, y, w, h, runs, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1, space_after=1)
    return b

def arrow_down(slide, cx, y, h=Inches(0.22)):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, cx - Pt(7), y, Pt(14), h)
    _solid(a, SUB)
    _shadow_off(a)

cx = Inches(4.7)        # 中列中心
colw = Inches(4.2)
colx = cx - colw / 2

y = Inches(2.2)
arch_box(s, colx, y, colw, Inches(0.62), "用户查询  User Query", INK, WHITE)
arrow_down(s, cx, y + Inches(0.66))
y += Inches(0.95)
arch_box(s, colx, y, colw, Inches(0.7), "智能路由  Router", BRAND, WHITE,
         "多维度决策 · 记忆增强路由")
arrow_down(s, cx, y + Inches(0.74))
y += Inches(1.03)
# Agent 群三个小框
aw = Inches(1.3)
agap = (colw - aw * 3) / 2
for i, name in enumerate(["技术 Agent", "创意 Agent", "记忆 Agent"]):
    arch_box(s, colx + i * (aw + agap), y, aw, Inches(0.66), name, ACCENT, WHITE,
             size=11.5)
arrow_down(s, cx, y + Inches(0.7))
y += Inches(0.99)
arch_box(s, colx, y, colw, Inches(0.7), "响应聚合  Aggregator", BRAND, WHITE,
         "加权合并 · 置信选择 · 冲突解决")
arrow_down(s, cx, y + Inches(0.74))
y += Inches(1.03)
arch_box(s, colx, y, colw, Inches(0.62), "最终响应  Final Response", INK, WHITE)

# 右侧支撑层
rx = Inches(9.5)
rw = Inches(3.0)
add_text(s, rx, Inches(2.2), rw, Pt(20), [("支撑层", 12, SUB, True)])
support = [
    ("长程记忆", "FAISS 向量库 · 记忆优化器", BRAND_D),
    ("分布式协调", "注册 · 负载均衡 · 资源锁", BRAND_D),
    ("故障恢复", "心跳 · 检查点 · 降级", BRAND_D),
]
sy = Inches(2.55)
for h, d, c in support:
    card = add_rect(s, rx, sy, rw, Inches(1.0), SOFT,
                    MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    add_rect(s, rx, sy + Inches(0.18), Pt(4), Inches(0.64), c)
    add_text(s, rx + Inches(0.22), sy + Inches(0.16), rw - Inches(0.4), Pt(26),
             [(h, 14, INK, True)])
    add_text(s, rx + Inches(0.22), sy + Inches(0.55), rw - Inches(0.4), Pt(24),
             [(d, 10.5, SUB, False)])
    sy += Inches(1.18)
footer(s, 3)

# ================================================================ 4-8. 五大核心功能
features = [
    ("CORE 01", "多Agent智能分流", "把对的问题交给对的 Agent", [
        ("多类型 Agent", "技术 / 创意 / 记忆 Agent，各司其职处理专属任务"),
        ("智能路由判断", "Router 分析查询类型与意图，匹配最合适 Agent"),
        ("多种路由策略", "直接路由 · 基于历史路由 · 记忆增强路由"),
        ("动态负载均衡", "结合系统负载实时分配，资源利用最优化"),
    ]),
    ("CORE 02", "长程记忆优化", "像人一样记住、压缩与遗忘", [
        ("向量化存储", "FAISS 高效相似度检索，支持百万级记忆库"),
        ("智能记忆压缩", "TF-IDF + 余弦相似度合并语义相近记忆"),
        ("重要性评分", "结合时间衰减、访问频率，遗忘低优先级记忆"),
        ("检索增强生成", "上下文感知检索 + 记忆融合，提升回答质量"),
    ]),
    ("CORE 03", "分布式协调", "多Agent的高可用编排", [
        ("注册与发现", "DistributedCoordinator 管理注册、注销与心跳"),
        ("多负载均衡策略", "轮询 / 随机 / 最少连接 / 性能 / 任务类型"),
        ("任务与消息", "任务分配、消息传递与资源锁机制"),
        ("水平扩展", "集群部署线性扩展，10 节点可达 850 请求/秒"),
    ]),
    ("CORE 04", "恢复管理", "故障来临时依旧稳健", [
        ("心跳监控", "实时监控 Agent 状态，快速发现异常"),
        ("检查点机制", "周期性保存状态，支持从检查点恢复"),
        ("故障隔离", "防止单点故障扩散到整个系统"),
        ("平滑降级", "部分故障时保持核心功能可用"),
    ]),
    ("CORE 05", "响应聚合", "把多路回答融成一个最优解", [
        ("多融合策略", "加权合并 · 置信度选择 · 简单合并"),
        ("冲突解决", "多数投票 / 置信加权 / 语义分析 / 源可靠性"),
        ("事实 vs 观点", "区分冲突类型，采用不同解决方案"),
        ("反馈闭环", "收集用户反馈，持续优化聚合与可靠性评分"),
    ]),
]

for idx, (kick, title, sub, items) in enumerate(features):
    s = new_slide()
    title_block(s, kick, title, sub)
    # 2x2 要点网格
    gw, gh = Inches(5.85), Inches(1.75)
    gx, gy = Inches(0.7), Inches(2.45)
    ggap = Inches(0.2)
    for j, (h, d) in enumerate(items):
        r, c = divmod(j, 2)
        x = gx + c * (gw + ggap)
        y = gy + r * (gh + ggap)
        add_rect(s, x, y, gw, gh, SOFT, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
        # 序号圆点
        num = add_rect(s, x + Inches(0.3), y + Inches(0.32), Inches(0.55),
                       Inches(0.55), BRAND, MSO_SHAPE.OVAL)
        add_text(s, x + Inches(0.3), y + Inches(0.32), Inches(0.55), Inches(0.55),
                 [(f"{j+1}", 16, WHITE, True)], align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + Inches(1.05), y + Inches(0.3), gw - Inches(1.35), Pt(28),
                 [(h, 16, INK, True)])
        add_text(s, x + Inches(1.05), y + Inches(0.78), gw - Inches(1.35),
                 Inches(0.85), [(d, 12, SUB, False)], line_spacing=1.3)
    footer(s, 4 + idx)

# ================================================================ 9. 技术亮点
s = new_slide()
title_block(s, "HIGHLIGHTS", "技术亮点",
            "融合分布式系统、NLP、向量检索与机器学习的工程实践")
hi = [
    ("智能冲突解决", "事实/观点冲突分类，语义相似度 + 矛盾检测，自适应选择解决策略"),
    ("自适应路由优化", "在线学习实时调整，A/B 测试评估，冷启动特殊处理"),
    ("分层记忆架构", "短/中/长期记忆分层，多维索引，时间衰减模型"),
    ("高质量工程", "模块化设计 · 全面类型提示 · 单元/集成/性能测试"),
    ("可配置可扩展", "配置驱动 · 插件架构 · 一致的 API 设计"),
    ("性能优化", "异步处理 · 多级缓存 · 请求批处理"),
]
gw, gh = Inches(3.85), Inches(1.55)
gx, gy = Inches(0.7), Inches(2.3)
gapx, gapy = Inches(0.18), Inches(0.18)
for j, (h, d) in enumerate(hi):
    r, c = divmod(j, 3)
    x = gx + c * (gw + gapx)
    y = gy + r * (gh + gapy)
    add_rect(s, x, y, gw, gh, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    # 边框感：用浅灰描边矩形覆层
    border = add_rect(s, x, y, gw, gh, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    border.fill.background()
    border.line.color.rgb = HAIR
    border.line.width = Pt(1)
    _shadow_off(border)
    add_rect(s, x + Inches(0.28), y + Inches(0.3), Pt(28), Pt(4), ACCENT)
    add_text(s, x + Inches(0.28), y + Inches(0.45), gw - Inches(0.55), Pt(26),
             [(h, 15, INK, True)])
    add_text(s, x + Inches(0.28), y + Inches(0.88), gw - Inches(0.55),
             Inches(0.6), [(d, 11, SUB, False)], line_spacing=1.25)
footer(s, 9)

# ================================================================ 10. 性能数据
s = new_slide()
title_block(s, "BENCHMARKS", "性能数据",
            "记忆增强路由准确率达 92.8%，综合优化后吞吐量翻倍")

# 关键指标大数字
metrics = [
    ("92.8%", "记忆增强路由准确率", "vs 直接路由 78.5%"),
    ("+102%", "系统吞吐量提升", "42 → 85 请求/秒"),
    ("-51.8%", "平均响应时间", "385ms → 186ms"),
    ("82.3%", "记忆存储压缩率", "相关性保持 87.5%"),
]
mw, mgap = Inches(2.95), Inches(0.18)
mx, my = Inches(0.7), Inches(2.35)
for i, (num, label, note) in enumerate(metrics):
    x = mx + i * (mw + mgap)
    add_rect(s, x, my, mw, Inches(1.85), SOFT, MSO_SHAPE.ROUNDED_RECTANGLE, 0.07)
    add_text(s, x, my + Inches(0.28), mw, Pt(46),
             [(num, 34, BRAND, True)], align=PP_ALIGN.CENTER)
    add_text(s, x, my + Inches(1.0), mw, Pt(24),
             [(label, 13, INK, True)], align=PP_ALIGN.CENTER)
    add_text(s, x, my + Inches(1.38), mw, Pt(22),
             [(note, 10.5, SUB, False)], align=PP_ALIGN.CENTER)

# 优化前后对比表
ty = Inches(4.55)
add_text(s, Inches(0.7), ty - Inches(0.05), Inches(8), Pt(22),
         [("优化前后对比", 14, INK, True)])
table_rows = [
    ("性能指标", "优化前", "优化后", "提升"),
    ("平均响应时间", "385.2 ms", "185.7 ms", "51.8%"),
    ("系统吞吐量", "42 请求/秒", "85 请求/秒", "102.4%"),
    ("内存使用", "4.2 GB", "2.8 GB", "33.3%"),
    ("路由准确率", "75.3%", "92.8%", "23.2%"),
]
rows, cols = len(table_rows), 4
tbl_w = Inches(11.93)
tx, tyy = Inches(0.7), ty + Inches(0.4)
gtbl = s.shapes.add_table(rows, cols, tx, tyy, tbl_w, Inches(1.7)).table
gtbl.columns[0].width = Inches(3.5)
for c in range(1, 4):
    gtbl.columns[c].width = Inches(2.81)
for r in range(rows):
    for c in range(cols):
        cell = gtbl.cell(r, c)
        cell.margin_left = Inches(0.15)
        cell.margin_top = Inches(0.04)
        cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = BRAND
        else:
            cell.fill.fore_color.rgb = WHITE if r % 2 else SOFT
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        run = para.add_run()
        run.text = table_rows[r][c]
        run.font.name = FONT
        run.font.size = Pt(11.5)
        run.font.bold = (r == 0) or (c == 3)
        if r == 0:
            run.font.color.rgb = WHITE
        elif c == 3:
            run.font.color.rgb = ACCENT
        else:
            run.font.color.rgb = INK
footer(s, 10)

# ================================================================ 11. 路线图
s = new_slide()
title_block(s, "ROADMAP", "后续规划",
            "从更智能的路由到多模态与知识图谱集成")
road = [
    ("近期", "更复杂的路由策略 · 记忆系统性能优化"),
    ("中期", "增强故障恢复智能化 · 完善分布式协调"),
    ("远期", "多模态 Agent · 自主学习 · Agent 辩论协商"),
    ("愿景", "知识图谱集成，构建领域知识增强推理"),
]
# 横向时间线
ly = Inches(3.6)
add_line(s, Inches(1.0), ly, Inches(11.3), HAIR, 2)
seg = Inches(11.3) / 4
for i, (phase, desc) in enumerate(road):
    cx = Inches(1.0) + seg * i + seg / 2
    dot = add_rect(s, cx - Pt(9), ly - Pt(9), Pt(18), Pt(18), BRAND, MSO_SHAPE.OVAL)
    add_text(s, cx - Inches(1.3), ly - Inches(1.0), Inches(2.6), Pt(26),
             [(phase, 16, BRAND, True)], align=PP_ALIGN.CENTER)
    add_text(s, cx - Inches(1.35), ly + Inches(0.35), Inches(2.7), Inches(1.4),
             [(desc, 12, SUB, False)], align=PP_ALIGN.CENTER, line_spacing=1.3)
footer(s, 11)

# ================================================================ 12. 结尾
s = new_slide()
add_rect(s, 0, 0, Pt(10), EMU_H, BRAND)
add_text(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(1.2),
         [("一个高性能、高可靠的", 30, INK, True),
          ("多Agent智能系统框架", 30, BRAND, True)], line_spacing=1.1)
add_line(s, Inches(0.95), Inches(4.5), Inches(2.2), BRAND, 2.5)
add_text(s, Inches(0.9), Inches(4.7), Inches(11), Pt(24),
         [("为复杂 AI 应用提供坚实的技术基础", 15, SUB, False)])
add_text(s, Inches(0.9), Inches(6.4), Inches(11), Pt(24),
         [("欢迎提交 Issue 与 Pull Request  ·  MIT License", 12, SUB, False)])
add_text(s, Inches(0.9), Inches(6.75), Inches(11), Pt(24),
         [("联系作者「林修」微信：LinXiu230624", 12, INK, True)])

import os
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "MultiAgent_Optimization.pptx")
prs.save(out)
print(f"已生成: {out}  共 {len(prs.slides._sldIdLst)} 页")
